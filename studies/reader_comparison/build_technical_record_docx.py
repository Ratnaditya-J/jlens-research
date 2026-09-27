"""Build the dated editable technical record from its reviewed Markdown source."""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def comparison_table(doc, block):
    rows = [[v.strip() for v in line.strip().strip('|').split('|')]
            for line in block.splitlines() if not re.fullmatch(r'[| :\-]+', line)]
    table = doc.add_table(rows=0, cols=5)
    table.autofit = False
    widths = [2.6, .75, .75, 1.65, 1.05]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if index == 0:
            props.append(OxmlElement('w:tblHeader'))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6EE' if index == 0 else 'FFFFFF')
            tcpr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + key), val)
                borders.append(border)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '100')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tcpr.append(margins)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(value)
            run.font.size = Pt(10.5)
            run.bold = index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def text_with_links(paragraph, text):
    pattern=r'https://[^\s]+'
    cursor=0
    for match in re.finditer(pattern,text):
        paragraph.add_run(text[cursor:match.start()])
        url=match.group().rstrip('.,')
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'244C6B');pr.append(c);r.append(pr)
        t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);paragraph._p.append(h)
        paragraph.add_run(match.group()[len(url):]);cursor=match.end()
    paragraph.add_run(text[cursor:])


def main():
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--out',type=Path,required=True);a=p.parse_args()
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8);section.left_margin=section.right_margin=Inches(.85)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',26),('Heading 1',16),('Heading 2',13)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(15 if name!='Title' else 0);style.paragraph_format.space_after=Pt(7);style.paragraph_format.keep_with_next=True
    for block in a.source.read_text().strip().split('\n\n'):
        if block.startswith('# '):doc.add_paragraph(block[2:].strip(),style='Title')
        elif block.startswith('## '):doc.add_paragraph(block[3:].strip(),style='Heading 1')
        elif block.startswith('|'):comparison_table(doc, block)
        else:
            paragraph=doc.add_paragraph();text_with_links(paragraph,block.replace('\n',' '));paragraph.paragraph_format.widow_control=True
    footer=section.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    doc.core_properties.title='Comparing activation probes J lens and Oracle Lens'
    doc.core_properties.subject='Technical methods and execution record'
    doc.core_properties.author='';doc.core_properties.keywords='J lens, Oracle Lens, activation probes, methods'
    for tree in [doc.styles.element, doc._element]:
        for border in list(tree.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    a.out.parent.mkdir(parents=True,exist_ok=True);doc.save(a.out)


if __name__=='__main__':main()
